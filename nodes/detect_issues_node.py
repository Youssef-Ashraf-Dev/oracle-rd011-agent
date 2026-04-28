"""
RD.011 Agent — Detect Issues node.

Uses a Reasoning model (DeepSeek R1) to perform a cross-file audit by
examining the *already-extracted* structured facts for contradictions,
ambiguities, and missing required fields.

IMPORTANT: raw_texts is intentionally NOT passed here.  Reasoning models
on free tiers have small context windows (~32k–128k tokens).  Feeding 50
pages of raw MoM directly causes 429 rate-limit errors.  The extraction
step already produced a compressed JSON of all key facts — use that instead.
"""

from __future__ import annotations

import json
import logging
import re
from pathlib import Path

from config import TEMPLATE_PATH, TaskType
from llm.retry import call_with_retry
from models.schemas import IssueReport

logger = logging.getLogger(__name__)

DETECT_ISSUES_PROMPT = """\
══════════════════════════════════════════════════════════
CRITICAL RULE — WHAT IS AND IS NOT A CONFLICT
══════════════════════════════════════════════════════════

A conflict ONLY exists when TWO documents BOTH contain an explicit,
contradictory value for the same fact.

These are NOT conflicts and must NOT be reported:
  • One document states a value; the other document is silent on the topic.
  • One side is "not specified", "not mentioned", "N/A", "TBD", or "unknown".
  • A process, feature, or requirement is only mentioned in one document.
  • Absence of information in one document is NOT a contradiction.

ONLY report a contradiction when you can state ALL THREE of:
  (1) "Document A explicitly says X"
  (2) "Document B explicitly says Y"
  (3) X ≠ Y  (they are genuinely incompatible, not just phrased differently)

If you cannot fill in both Document A and Document B with explicit,
incompatible values — it is NOT a conflict. Do not report it.

══════════════════════════════════════════════════════════

You are a senior Oracle Finance Cloud quality auditor. You have been given:
1. The structured extraction results — a compressed JSON of all key facts
   distilled from the input files (MoMs, Scope documents, Questionnaires).
2. The proposed document plan.
3. The RD.011 template placeholder fields (if available).

IMPORTANT - DO NOT REPORT TEMPLATE/STRUCTURE SECTIONS AS "MISSING":
The pipeline will ALWAYS generate the following standard RD.011 sections from the template
and the generator output. They are NOT missing even if the MoM does not mention them:
- Document Control, Change Record, Reviewers
- Module chapter structure, process outline tables
- Process narratives, step catalogs, diagrams, journal entries, key requirements
- Open/Closed issues sections

Your "missing_required_fields" must ONLY contain missing CLIENT-SPECIFIC INPUTS that block
accurate generation (e.g., approval limits matrix, payment terms, COA segment lengths),
not "sections that should exist in the document".

The extraction step has already identified raw conflicts in the
"conflicts_between_documents" field. Your job is to go further:

A. VALIDATE the listed conflicts — are they genuine? Are both values really \
   incompatible, or is one a superset/clarification of the other?

B. IDENTIFY any ADDITIONAL cross-fact contradictions that the extraction may \
   have missed by examining these areas:

   1. **COA Structure Consistency**: Does the Chart of Accounts structure match \
   across GL-related facts? Are segment names, counts, and hierarchies consistent?

   2. **Approval Level Conflicts**: Are there conflicting approval levels or limits \
   across AP and Procurement data? Do approval hierarchies match across modules?

   3. **Currency Consistency**: Are functional currencies consistent across all \
   module facts? Check for discrepancies in currency codes, multi-currency setup, \
   and reporting currency definitions.

   4. **Legal Entity Consistency**: Is the number and naming of legal entities \
   consistent throughout the extracted facts?

   5. **Matching / Process Rules**: Do any process rules (matching method, \
   netting, intercompany, payment terms) conflict between modules or documents?

   6. **General Contradictions**: Any other data points that contradict between \
   earlier MoM extractions and later scope/summary extractions.

C. IDENTIFY missing required fields: any data needed for the RD.011 document \
   that is not adequately covered in the extracted facts. Use ONLY the template \
   placeholders as a checklist for required content (do not list generic \
   template section headings).

D. LIST ambiguities per module section: requirements that are unclear or need \
   consultant clarification before content generation begins.

When reporting contradictions, reference the source documents and recommend \
a specific resolution — do not just describe the conflict, tell the consultant \
which value to use and why.

Return ONLY valid JSON with this structure:
{
  "contradictions": [
        "Approval matrix conflicts: Level-2 approver is Finance Manager in MOM_Day2.docx vs Chief Accountant in Scope_v2.docx. RESOLUTION: Use Scope_v2.docx value because it is the latest approved baseline."
  ],
  "ambiguities_by_section": {
    "AP": ["Ambiguity 1 with resolution suggestion...", "Ambiguity 2..."],
    "GL": ["Ambiguity 1..."]
  },
  "missing_required_fields": ["Missing field 1 — impact on document...", ...]
}
"""



def _extract_template_requirements() -> dict:
    """
    Extract placeholder-like fields from the RD.011 template.
    """
    template_path = Path(TEMPLATE_PATH)
    if not template_path.exists():
        logger.warning("Template not found at %s", template_path)
        return {"placeholders": []}

    try:
        from docx import Document
    except Exception as exc:
        logger.warning("python-docx not available: %s", exc)
        return {"placeholders": []}

    try:
        doc = Document(str(template_path))
        placeholders = []
        seen_placeholders = set()
        placeholder_patterns = [
            re.compile(r"<<[^>]+>>"),
            re.compile(r"\{\{[^}]+\}\}"),
        ]

        for para in doc.paragraphs:
            text = (para.text or "").strip()
            if not text:
                continue
            for pattern in placeholder_patterns:
                for match in pattern.findall(text):
                    if match not in seen_placeholders:
                        placeholders.append(match)
                        seen_placeholders.add(match)

        return {"placeholders": placeholders}
    except Exception as exc:
        logger.warning("Failed to parse template: %s", exc)
        return {"placeholders": []}


def detect_issues_node(state: dict) -> dict:
    """
    Perform cross-fact audit to detect contradictions, ambiguities,
    and missing required fields within the extracted structured data.

    Receives extraction_result and document_plan only — raw_texts is
    deliberately excluded to keep the prompt within reasoning-model
    context window limits.
    """
    extraction_data = state.get("extraction_result", {})
    plan_data = state.get("document_plan", {})
    errors = list(state.get("errors", []))

    extraction_json = json.dumps(extraction_data, indent=2, ensure_ascii=False)
    plan_json = json.dumps(plan_data, indent=2, ensure_ascii=False)
    template_requirements = _extract_template_requirements()
    template_json = json.dumps(template_requirements, indent=2, ensure_ascii=False)

    full_prompt = f"""{DETECT_ISSUES_PROMPT}

## Extraction Results (compressed facts from all input files)

{extraction_json}

## Document Plan

{plan_json}

## Template Requirements (placeholders only)

{template_json}

Perform the cross-fact audit now and return the JSON IssueReport."""

    try:
        report = call_with_retry(TaskType.REASONING, full_prompt, IssueReport)
        logger.info(
            "Issue detection complete: %d contradictions, %d sections with ambiguities",
            len(report.contradictions),
            len(report.ambiguities_by_section),
        )
        return {
            "issue_report": report.model_dump(),
            "last_completed_node": "detect_issues",
        }
    except RuntimeError as exc:
        error_msg = f"Issue detection failed: {exc}"
        logger.error(error_msg)
        errors.append(error_msg)
        # Return empty report so the pipeline can continue
        empty_report = IssueReport(
            contradictions=[],
            ambiguities_by_section={},
            missing_required_fields=[],
        )
        return {
            "issue_report": empty_report.model_dump(),
            "errors": errors,
            "last_completed_node": "detect_issues",
        }
